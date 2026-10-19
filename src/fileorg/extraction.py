"""Extract text and embedded metadata from PDF, DOCX and TXT files.

Design notes
------------
* Extraction never raises: problems are reported in ``ExtractionResult`` so a
  single broken file cannot stop a scan of thousands of files.
* Text is capped at ``max_chars`` to keep memory and run time predictable.
* Scanned PDFs (images only) yield no text and get status ``empty``. OCR is
  out of scope for version 1.
"""
from __future__ import annotations

import codecs
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional

DEFAULT_MAX_CHARS = 100_000


@dataclass
class ExtractionResult:
    text: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)
    status: str = "ok"  # ok | empty | error | unsupported
    error: Optional[str] = None


# --------------------------------------------------------------------- TXT
def _decode(data: bytes, truncated: bool) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    decoder = codecs.getincrementaldecoder("utf-8-sig")()
    try:
        # final=False tolerates a multi-byte character cut by truncation
        return decoder.decode(data, final=not truncated)
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1250")  # Central European legacy text files
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def _extract_txt(path: Path, max_chars: int, max_pages: Optional[int]) -> tuple[str, dict]:
    max_bytes = max_chars * 4
    with path.open("rb") as fh:
        data = fh.read(max_bytes + 1)
    truncated = len(data) > max_bytes
    return _decode(data[:max_bytes], truncated), {}


# --------------------------------------------------------------------- PDF
def _extract_pdf(path: Path, max_chars: int, max_pages: Optional[int]) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted and int(reader.decrypt("")) == 0:
        raise ValueError("PDF is password protected")

    meta: dict[str, Any] = {"pages": len(reader.pages)}
    info = reader.metadata
    if info:
        for key in ("title", "author", "subject", "creator", "producer"):
            value = getattr(info, key, None)
            if value:
                meta[key] = str(value)
        try:
            created = info.creation_date
            if created:
                meta["created"] = created.isoformat()
        except Exception:  # malformed date strings are common
            pass

    parts: list[str] = []
    total = 0
    for index, page in enumerate(reader.pages):
        if max_pages is not None and index >= max_pages:
            break
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        parts.append(page_text)
        total += len(page_text)
        if total >= max_chars:
            break
    return "\n".join(parts), meta


# -------------------------------------------------------------------- DOCX
def _extract_docx(path: Path, max_chars: int, max_pages: Optional[int]) -> tuple[str, dict]:
    from docx import Document

    doc = Document(str(path))
    props = doc.core_properties
    meta: dict[str, Any] = {}
    for key in ("title", "author", "subject", "keywords", "category", "last_modified_by"):
        value = getattr(props, key, None)
        if value:
            meta[key] = str(value)
    for key in ("created", "modified"):
        value = getattr(props, key, None)
        if value:
            meta[key] = value.isoformat()

    parts: list[str] = []
    total = 0
    # Paragraphs first, then tables. (Their original interleaving is not kept;
    # that is fine for keyword and bag-of-words features.)
    for paragraph in doc.paragraphs:
        parts.append(paragraph.text)
        total += len(paragraph.text)
        if total >= max_chars:
            return "\n".join(parts), meta
    for table in doc.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            parts.append(line)
            total += len(line)
            if total >= max_chars:
                return "\n".join(parts), meta
    return "\n".join(parts), meta


_EXTRACTORS: dict[str, Callable[[Path, int, Optional[int]], tuple[str, dict]]] = {
    ".txt": _extract_txt,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def _normalize(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract(
    path: str | Path,
    max_chars: int = DEFAULT_MAX_CHARS,
    max_pages: Optional[int] = None,
) -> ExtractionResult:
    path = Path(path)
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        return ExtractionResult(status="unsupported")
    try:
        raw_text, meta = extractor(path, max_chars, max_pages)
    except Exception as exc:  # noqa: BLE001 - report, never crash a scan
        return ExtractionResult(status="error", error=f"{type(exc).__name__}: {exc}")

    text = _normalize(raw_text)[:max_chars]
    return ExtractionResult(text=text, metadata=meta, status="ok" if text else "empty")
